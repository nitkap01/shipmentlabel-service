import io
import time

from openpyxl import Workbook, load_workbook

from app.bulk.template import TEMPLATE_HEADERS
from tests.conftest import make_ship_success


def _upload_file_bytes(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(TEMPLATE_HEADERS)
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _wait_for_terminal_status(client, run_id: int, timeout_seconds: float = 15) -> dict:
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        run = client.get(f"/api/bulk/runs/{run_id}").json()
        if run["status"] in ("completed", "completed_with_errors", "failed"):
            return run
        time.sleep(0.3)
    raise AssertionError("Bulk run did not reach a terminal status in time")


def test_full_bulk_pipeline_upload_to_report_and_zip(logged_in_client, monkeypatch, fake_png_base64):
    calls = {"n": 0}

    async def fake_ship(environment, body):
        calls["n"] += 1
        return make_ship_success(fake_png_base64, tracking=f"T{calls['n']}", unique_ref=f"U{calls['n']}")

    monkeypatch.setattr("app.labels_service.epg_client.ship", fake_ship)

    file_bytes = _upload_file_bytes(
        [
            ["Row One", None, "1 First St", None, "Newark", "NJ", "07102", None, None, 4, "oz", None, None, None, None, 5.0, "R1", None],
            ["Row Two", None, "2 Second St", None, "Trenton", "NJ", "08619", None, None, 8, "oz", None, None, None, None, 7.5, "R2", None],
            ["Bad Row", None, None, None, "Trenton", "NJ", "08619", None, None, 8, "oz", None, None, None, None, 7.5, "R3", None],
        ]
    )

    upload_resp = logged_in_client.post(
        "/api/bulk/uploads",
        files={"file": ("test.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert upload_resp.status_code == 200
    upload_body = upload_resp.json()
    assert upload_body["run"]["total_rows"] == 3
    assert upload_body["run"]["valid_rows"] == 2
    assert upload_body["row_errors"] == [{"row_number": 3, "error": "Address Line 1 is required"}]

    run_id = upload_body["run"]["id"]

    start_resp = logged_in_client.post(f"/api/bulk/runs/{run_id}/start")
    assert start_resp.status_code == 200
    assert start_resp.json()["status"] == "queued"

    final_run = _wait_for_terminal_status(logged_in_client, run_id)
    assert final_run["status"] == "completed_with_errors"
    assert final_run["success_count"] == 2
    assert final_run["failure_count"] == 1
    assert calls["n"] == 2  # the invalid row never reached EPG

    report_resp = logged_in_client.get(f"/api/bulk/runs/{run_id}/report.xlsx")
    assert report_resp.status_code == 200
    wb = load_workbook(io.BytesIO(report_resp.content))
    assert wb.sheetnames == ["All Rows", "Failed Rows"]

    zip_resp = logged_in_client.get(f"/api/bulk/runs/{run_id}/labels.zip")
    assert zip_resp.status_code == 200
    import zipfile

    with zipfile.ZipFile(io.BytesIO(zip_resp.content)) as zf:
        assert len(zf.namelist()) == 2


def test_upload_rejects_file_with_wrong_headers(logged_in_client):
    wb = Workbook()
    ws = wb.active
    ws.append(["Recipient Name", "Something Else"])
    ws.append(["Jane", "x"])
    out = io.BytesIO()
    wb.save(out)

    resp = logged_in_client.post(
        "/api/bulk/uploads",
        files={"file": ("bad.xlsx", out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 400
