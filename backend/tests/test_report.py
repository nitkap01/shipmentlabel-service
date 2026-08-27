import io
import zipfile
from types import SimpleNamespace

from openpyxl import load_workbook

from app.bulk.report import build_labels_zip, build_report_xlsx
from app.bulk.template import TEMPLATE_HEADERS
from app.storage import write_bytes_under_root


def _row(row_number, status, raw, label_id=None, validation_error=None, epg_error_message=None):
    return SimpleNamespace(
        row_number=row_number,
        status=status,
        raw=raw,
        label_id=label_id,
        validation_error=validation_error,
        epg_error_message=epg_error_message,
    )


def _raw_row(name: str) -> dict:
    return {h: None for h in TEMPLATE_HEADERS} | {
        "Recipient Name": name,
        "Address Line 1": "123 Main St",
        "City": "Newark",
        "State": "NJ",
        "Postal Code": "07102",
        "Weight": 4,
        "Weight Unit": "oz",
        "Declared Value": 10.0,
    }


def test_report_has_both_sheets_with_one_row_per_data_row():
    rows = [
        _row(1, "success", _raw_row("Alice"), label_id=1),
        _row(2, "failed", _raw_row("Bob"), epg_error_message="EPG rejected it"),
        _row(3, "invalid", _raw_row("Carol"), validation_error="Bad state"),
    ]
    content = build_report_xlsx(rows, labels_by_id={})
    wb = load_workbook(io.BytesIO(content))

    assert wb.sheetnames == ["All Rows", "Failed Rows"]
    all_rows = list(wb["All Rows"].iter_rows(values_only=True))
    assert len(all_rows) == 4  # header + 3 data rows
    assert all_rows[1][:2] == (1, "success")
    assert all_rows[2][:3] == (2, "failed", "EPG rejected it")
    assert all_rows[3][:3] == (3, "invalid", "Bad state")


def test_failed_rows_sheet_headers_are_byte_identical_to_template():
    rows = [_row(1, "failed", _raw_row("Bob"), epg_error_message="oops")]
    content = build_report_xlsx(rows, labels_by_id={})
    wb = load_workbook(io.BytesIO(content))

    failed_header = list(wb["Failed Rows"].iter_rows(values_only=True))[0]
    assert list(failed_header) == TEMPLATE_HEADERS


def test_failed_rows_sheet_only_contains_failed_and_invalid_rows():
    rows = [
        _row(1, "success", _raw_row("Alice")),
        _row(2, "failed", _raw_row("Bob"), epg_error_message="err"),
        _row(3, "invalid", _raw_row("Carol"), validation_error="bad"),
        _row(4, "skipped", _raw_row("Dave")),
    ]
    content = build_report_xlsx(rows, labels_by_id={})
    wb = load_workbook(io.BytesIO(content))
    failed_data_rows = list(wb["Failed Rows"].iter_rows(values_only=True))[1:]
    names = [r[0] for r in failed_data_rows]
    assert names == ["Bob", "Carol"]


def test_zip_contains_exactly_successful_rows_pdfs(monkeypatch, tmp_path):
    monkeypatch.setattr("app.storage.settings.label_storage_root", str(tmp_path))

    path1, _ = write_bytes_under_root("labels/2026/08", "one.pdf", b"pdf-one")
    path2, _ = write_bytes_under_root("labels/2026/08", "two.pdf", b"pdf-two")

    labels = [
        SimpleNamespace(pdf_path=path1),
        SimpleNamespace(pdf_path=path2),
        SimpleNamespace(pdf_path=None),  # a failed/pending label with no PDF must be skipped
    ]

    content = build_labels_zip(labels)
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        names = sorted(zf.namelist())
        assert names == ["one.pdf", "two.pdf"]
        assert zf.read("one.pdf") == b"pdf-one"
        assert zf.read("two.pdf") == b"pdf-two"
