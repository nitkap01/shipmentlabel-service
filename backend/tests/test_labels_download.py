import io
import zipfile

from tests.conftest import make_ship_success

VALID_LABEL_PAYLOAD = {
    "recipient_name": "Jane Doe",
    "recipient_address1": "123 Main St",
    "recipient_city": "Newark",
    "recipient_state": "NJ",
    "recipient_postal_code": "07102",
    "weight_value": 8,
    "weight_unit": "oz",
    "declared_value": 10.00,
}


def _create_label(client, monkeypatch, tracking: str, unique_ref: str, fake_png_base64: str) -> int:
    async def fake_ship(environment, body):
        return make_ship_success(fake_png_base64, tracking=tracking, unique_ref=unique_ref)

    monkeypatch.setattr("app.labels_service.epg_client.ship", fake_ship)
    resp = client.post("/api/labels", json=VALID_LABEL_PAYLOAD)
    assert resp.status_code == 200
    return resp.json()["id"]


def test_download_zip_contains_selected_labels_only(logged_in_client, monkeypatch, fake_png_base64):
    id1 = _create_label(logged_in_client, monkeypatch, "T1", "U1", fake_png_base64)
    id2 = _create_label(logged_in_client, monkeypatch, "T2", "U2", fake_png_base64)
    _create_label(logged_in_client, monkeypatch, "T3", "U3", fake_png_base64)

    resp = logged_in_client.post("/api/labels/download", json={"label_ids": [id1, id2]})
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/zip"

    with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
        assert len(zf.namelist()) == 2


def test_download_with_no_ids_returns_400(logged_in_client):
    resp = logged_in_client.post("/api/labels/download", json={"label_ids": []})
    assert resp.status_code == 400


def test_download_with_unknown_ids_returns_404(logged_in_client):
    resp = logged_in_client.post("/api/labels/download", json={"label_ids": [999999]})
    assert resp.status_code == 404


def test_report_xlsx_has_one_row_per_label_plus_header(logged_in_client, monkeypatch, fake_png_base64):
    from openpyxl import load_workbook

    _create_label(logged_in_client, monkeypatch, "R1", "U1", fake_png_base64)
    _create_label(logged_in_client, monkeypatch, "R2", "U2", fake_png_base64)

    resp = logged_in_client.get("/api/labels/report.xlsx")
    assert resp.status_code == 200

    wb = load_workbook(io.BytesIO(resp.content))
    sheet = wb["Labels"]
    assert sheet.max_row == 3
    tracking_values = {sheet.cell(row=r, column=6).value for r in (2, 3)}
    assert tracking_values == {"R1", "R2"}


def test_report_xlsx_respects_date_range(logged_in_client, monkeypatch, fake_png_base64):
    from openpyxl import load_workbook

    _create_label(logged_in_client, monkeypatch, "R1", "U1", fake_png_base64)

    resp = logged_in_client.get("/api/labels/report.xlsx", params={"from_date": "2000-01-01", "to_date": "2000-01-02"})
    assert resp.status_code == 200

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb["Labels"].max_row == 1


def test_search_matches_tracking_number(logged_in_client, monkeypatch, fake_png_base64):
    _create_label(logged_in_client, monkeypatch, "EPGUNIQUE555", "U1", fake_png_base64)
    _create_label(logged_in_client, monkeypatch, "OTHER999", "U2", fake_png_base64)

    resp = logged_in_client.get("/api/labels", params={"q": "epgunique555"})
    assert resp.status_code == 200
    items = resp.json()["items"]
    assert len(items) == 1
    assert items[0]["tracking_number"] == "EPGUNIQUE555"
