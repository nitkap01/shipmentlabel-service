"""SHIP-2 (record saved before buying), SHIP-3 (no double purchase), SHIP-5 (void only when ePost confirms)."""
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text

from app.bulk.report import build_report_xlsx
from app.bulk.runner import claim_next_run, process_run
from app.db import SessionLocal
from app.epg import mapping
from app.epg.client import EPGError, EPGTimeoutError
from app.models import BulkRunRow, Label
from tests.conftest import make_ship_success
from tests.test_bulk_runner import _make_queued_run

pytestmark = pytest.mark.duplicate_guard

PAYLOAD = {
    "recipient_name": "Jane Doe",
    "recipient_address1": "123 Main St",
    "recipient_city": "Newark",
    "recipient_state": "NJ",
    "recipient_postal_code": "07102",
    "weight_value": 8,
    "weight_unit": "oz",
    "declared_value": 10.00,
    "reference1": "ORDER-1",
}


@pytest.fixture
def ship_calls(monkeypatch, fake_png_base64):
    calls = []

    async def fake_ship(environment, body):
        calls.append(environment)
        return make_ship_success(fake_png_base64, tracking=f"EPGT{len(calls)}", unique_ref=f"u{len(calls)}")

    monkeypatch.setattr("app.labels_service.epg_client.ship", fake_ship)
    return calls


# ---- SHIP-2 -------------------------------------------------------------------------------------------------------

async def test_label_row_is_committed_before_epost_is_called(monkeypatch):
    """The row must already be in the database (seen from ANOTHER session) at the moment ePost is asked to buy."""
    seen = {}

    async def fake_ship(environment, body):
        async with SessionLocal() as other:
            rows = (await other.execute(select(Label))).scalars().all()
            seen["rows"] = [(r.status, r.epg_request_json is not None) for r in rows]
        raise RuntimeError("process crashed right after the purchase")

    monkeypatch.setattr("app.labels_service.epg_client.ship", fake_ship)
    from app.labels_service import create_label, get_settings_row

    async with SessionLocal() as db:
        settings_row = await get_settings_row(db)
        fields = dict(PAYLOAD, service_code="EP05", dimension_unit=None, length_value=None,
                      weight_value=Decimal("8"), declared_value=Decimal("10.00"))  # as the API's schema delivers them
        with pytest.raises(RuntimeError):
            await create_label(db, fields=fields, source="single", bulk_run_id=None, settings_row=settings_row)
        await db.rollback()

    assert seen["rows"] == [("pending", True)]
    async with SessionLocal() as db:  # and it survives the crash as "pending" = needs checking
        rows = (await db.execute(select(Label))).scalars().all()
        assert [r.status for r in rows] == ["pending"]


# ---- SHIP-3 -------------------------------------------------------------------------------------------------------

def test_identical_single_label_within_minutes_is_refused(logged_in_client, ship_calls):
    first = logged_in_client.post("/api/labels", json=PAYLOAD)
    assert first.status_code == 200 and first.json()["status"] == "created"
    again = logged_in_client.post("/api/labels", json=PAYLOAD)
    assert again.status_code == 409
    assert "identical label" in again.json()["detail"]
    assert len(ship_calls) == 1  # nothing bought the second time


def test_different_reference_is_allowed(logged_in_client, ship_calls):
    assert logged_in_client.post("/api/labels", json=PAYLOAD).status_code == 200
    other = dict(PAYLOAD, reference1="ORDER-2")
    assert logged_in_client.post("/api/labels", json=other).status_code == 200
    assert len(ship_calls) == 2


def test_old_identical_label_does_not_block(logged_in_client, ship_calls):
    assert logged_in_client.post("/api/labels", json=PAYLOAD).status_code == 200

    async def age():
        async with SessionLocal() as db:
            await db.execute(text("UPDATE labels SET created_at = created_at - interval '20 minutes'"))
            await db.commit()

    import asyncio
    asyncio.run(age())
    assert logged_in_client.post("/api/labels", json=PAYLOAD).status_code == 200


def test_uncertain_label_blocks_an_identical_one_for_a_day(logged_in_client, monkeypatch, fake_png_base64):
    async def timeout_ship(environment, body):
        raise EPGTimeoutError("timed out")

    monkeypatch.setattr("app.labels_service.epg_client.ship", timeout_ship)
    first = logged_in_client.post("/api/labels", json=PAYLOAD)
    assert first.json()["status"] == "pending"

    async def age():
        async with SessionLocal() as db:
            await db.execute(text("UPDATE labels SET created_at = created_at - interval '3 hours'"))
            await db.commit()

    import asyncio
    asyncio.run(age())
    again = logged_in_client.post("/api/labels", json=PAYLOAD)
    assert again.status_code == 409
    assert "needs checking" in again.json()["detail"]


async def test_bulk_reupload_of_a_bought_row_is_not_bought_again(monkeypatch, fake_png_base64):
    calls = []

    async def fake_ship(environment, body):
        calls.append(1)
        return make_ship_success(fake_png_base64, tracking=f"T{len(calls)}", unique_ref=f"u{len(calls)}")

    monkeypatch.setattr("app.labels_service.epg_client.ship", fake_ship)
    monkeypatch.setattr("app.bulk.runner.SEQUENTIAL_DELAY_SECONDS", 0)
    async with SessionLocal() as db:
        await _make_queued_run(db, 1)
        await process_run(db, await claim_next_run(db))
        await _make_queued_run(db, 1)  # the same row uploaded again a minute later
        run2 = await claim_next_run(db)
        await process_run(db, run2)
        rows = (await db.execute(select(BulkRunRow).where(BulkRunRow.bulk_run_id == run2.id))).scalars().all()
        assert rows[0].status == "duplicate"
        assert rows[0].label_id is not None
    assert len(calls) == 1


async def test_identical_rows_inside_one_bulk_file_are_all_bought(monkeypatch, fake_png_base64):
    calls = []

    async def fake_ship(environment, body):
        calls.append(1)
        return make_ship_success(fake_png_base64, tracking=f"T{len(calls)}", unique_ref=f"u{len(calls)}")

    monkeypatch.setattr("app.labels_service.epg_client.ship", fake_ship)
    monkeypatch.setattr("app.bulk.runner.SEQUENTIAL_DELAY_SECONDS", 0)
    async with SessionLocal() as db:
        run = await _make_queued_run(db, 2)
        for r in (await db.execute(select(BulkRunRow).where(BulkRunRow.bulk_run_id == run.id))).scalars():
            r.raw = dict(r.raw, **{"Recipient Name": "Same Person", "Reference / Order #": "SAME"})
        await db.commit()
        await process_run(db, await claim_next_run(db))
    assert len(calls) == 2  # two boxes to one customer in one file is normal


def test_report_keeps_uncertain_and_duplicate_rows_out_of_failed_rows():
    from types import SimpleNamespace as NS
    raw = {"Recipient Name": "X"}
    rows = [NS(row_number=1, status="failed", raw=raw, label_id=None, validation_error=None, epg_error_message="bad zip"),
            NS(row_number=2, status="needs_checking", raw=raw, label_id=7, validation_error=None, epg_error_message="timed out"),
            NS(row_number=3, status="duplicate", raw=raw, label_id=8, validation_error=None, epg_error_message="identical")]
    wb = load_workbook(io.BytesIO(build_report_xlsx(rows, {})))
    assert wb["Failed Rows"].max_row == 2  # header + the one truly failed row
    check = wb["Needs checking - do NOT re-upload"]
    assert [c.value for c in check["A"]][1:] == [2, 3]


# ---- SHIP-5 -------------------------------------------------------------------------------------------------------

def _created(logged_in_client, ship_calls):
    resp = logged_in_client.post("/api/labels", json=PAYLOAD)
    assert resp.status_code == 200
    return resp.json()


def test_void_not_confirmed_by_epost_is_not_marked_voided(logged_in_client, ship_calls, monkeypatch):
    label = _created(logged_in_client, ship_calls)

    async def fake_void(environment, uid):
        return [{"package": {"id": uid}, "success": False, "secondaryMessage": "Package already manifested"}]

    monkeypatch.setattr("app.routers.labels.epg_client.void", fake_void)
    resp = logged_in_client.post(f"/api/labels/{label['id']}/void")
    assert resp.status_code == 502
    assert "already manifested" in resp.json()["detail"]
    body = logged_in_client.get(f"/api/labels/{label['id']}").json()
    assert body["status"] == "created"
    assert body["void_error"] == "Package already manifested"


def test_void_confirmed_by_epost_list_reply(logged_in_client, ship_calls, monkeypatch):
    label = _created(logged_in_client, ship_calls)

    async def fake_void(environment, uid):  # the real sandbox shape (09_10_2026)
        return [{"package": {"id": uid, "status": "Closed"}, "success": True, "secondaryMessage": "Package voided"}]

    monkeypatch.setattr("app.routers.labels.epg_client.void", fake_void)
    resp = logged_in_client.post(f"/api/labels/{label['id']}/void")
    assert resp.status_code == 200 and resp.json()["status"] == "voided"


def test_void_uses_the_labels_own_environment(logged_in_client, ship_calls, monkeypatch):
    label = _created(logged_in_client, ship_calls)  # bought in sandbox

    async def switch():
        async with SessionLocal() as db:
            await db.execute(text("UPDATE app_settings SET epg_environment='production' WHERE id=1"))
            await db.commit()

    import asyncio
    asyncio.run(switch())
    used = []

    async def fake_void(environment, uid):
        used.append(environment)
        return [{"success": True}]

    monkeypatch.setattr("app.routers.labels.epg_client.void", fake_void)
    assert logged_in_client.post(f"/api/labels/{label['id']}/void").status_code == 200
    assert used == ["sandbox"]


def test_void_404_message_is_shown(logged_in_client, ship_calls, monkeypatch):
    label = _created(logged_in_client, ship_calls)

    async def fake_void(environment, uid):
        raise EPGError("This requested shipment does not exist or has already been voided.", status_code=404)

    monkeypatch.setattr("app.routers.labels.epg_client.void", fake_void)
    resp = logged_in_client.post(f"/api/labels/{label['id']}/void")
    assert resp.status_code == 502
    assert "already been voided" in resp.json()["detail"]


@pytest.mark.parametrize("reply,ok", [
    ([{"success": True}], True),
    ({"success": True}, True),
    ([{"success": True}, {"success": False}], False),
    ([], False),
    ({"message": "x"}, False),
    (None, False),
])
def test_void_result(reply, ok):
    assert mapping.void_result(reply)[0] is ok
