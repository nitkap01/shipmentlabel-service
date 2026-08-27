"""Parse and validate an uploaded bulk Excel file.

Structural check (headers) fails the whole file. Per-row validation never
calls EPG (D15) — it only checks shape, so bad rows are caught before any
money is spent.
"""

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.bulk.template import NORMALIZED_HEADERS, TEMPLATE_HEADERS, normalize_header
from app.units import DIMENSION_UNITS, WEIGHT_UNITS
from app.us_states import US_STATE_CODES

POSTAL_CODE_RE = re.compile(r"^\d{5}(-\d{4})?$")
VALID_SERVICE_CODES = {"EP03", "EP05"}


class HeaderMismatchError(ValueError):
    def __init__(self, missing: list[str], unknown: list[str]):
        self.missing = missing
        self.unknown = unknown
        parts = []
        if missing:
            parts.append(f"missing columns: {', '.join(missing)}")
        if unknown:
            parts.append(f"unknown columns: {', '.join(unknown)}")
        super().__init__("; ".join(parts) or "Header mismatch")


@dataclass
class ParsedRow:
    row_number: int
    raw: dict
    valid: bool
    error: str | None = None
    label_fields: dict = field(default_factory=dict)


def _cell_str(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _decimal(value) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).strip())
    except InvalidOperation:
        return None


def check_headers(header_row: tuple) -> None:
    normalized = [normalize_header(h) for h in header_row if h is not None]
    missing = [
        TEMPLATE_HEADERS[i]
        for i, expected in enumerate(NORMALIZED_HEADERS)
        if expected not in normalized
    ]
    unknown = [h for h in header_row if h is not None and normalize_header(h) not in NORMALIZED_HEADERS]
    if missing or unknown:
        raise HeaderMismatchError(missing, [str(u) for u in unknown])


def parse_workbook(
    file_bytes: bytes,
    weight_unit_override: str | None,
    dimension_unit_override: str | None,
    default_service_code: str,
    max_rows: int,
) -> list[ParsedRow]:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        raise HeaderMismatchError(TEMPLATE_HEADERS, [])
    check_headers(header_row)

    col_index = {normalize_header(h): i for i, h in enumerate(header_row) if h is not None}

    def get(row: tuple, name: str):
        idx = col_index.get(normalize_header(name))
        return row[idx] if idx is not None and idx < len(row) else None

    results: list[ParsedRow] = []
    for i, row in enumerate(rows_iter, start=1):
        if row is None or all(c is None for c in row):
            continue
        if i > max_rows:
            results.append(
                ParsedRow(row_number=i, raw={}, valid=False, error=f"Exceeds max row cap of {max_rows}")
            )
            continue
        results.append(_validate_row(i, row, get, weight_unit_override, dimension_unit_override, default_service_code))

    return results


def validate_raw_row(
    row_number: int,
    raw: dict,
    weight_unit_override: str | None,
    dimension_unit_override: str | None,
    default_service_code: str,
) -> ParsedRow:
    """Re-derive a ParsedRow from a stored `raw` dict (bulk_run_rows.raw).

    Used by the bulk runner to reconstruct label fields for a row that was
    already validated at upload time, without a second jsonb column.
    """
    return _validate_row(
        row_number,
        None,
        lambda _row, name: raw.get(name),
        weight_unit_override,
        dimension_unit_override,
        default_service_code,
    )


def _validate_row(
    row_number: int,
    row: tuple,
    get,
    weight_unit_override: str | None,
    dimension_unit_override: str | None,
    default_service_code: str,
) -> ParsedRow:
    raw = {h: get(row, h) for h in TEMPLATE_HEADERS}

    name = _cell_str(get(row, "Recipient Name"))
    address1 = _cell_str(get(row, "Address Line 1"))
    city = _cell_str(get(row, "City"))
    state = _cell_str(get(row, "State"))
    postal_code = _cell_str(get(row, "Postal Code"))
    weight_raw = _decimal(get(row, "Weight"))
    weight_unit_col = _cell_str(get(row, "Weight Unit"))
    declared_value = _decimal(get(row, "Declared Value"))
    service_code = _cell_str(get(row, "Service Code"))

    errors = []
    if not name:
        errors.append("Recipient Name is required")
    if not address1:
        errors.append("Address Line 1 is required")
    if not city:
        errors.append("City is required")
    if not state or state.upper() not in US_STATE_CODES:
        errors.append("State must be a valid 2-letter US code")
    if not postal_code or not POSTAL_CODE_RE.match(postal_code):
        errors.append("Postal Code must be NNNNN or NNNNN-NNNN")
    if weight_raw is None or weight_raw <= 0:
        errors.append("Weight must be a positive number")
    if declared_value is None or declared_value < 0:
        errors.append("Declared Value must be a non-negative number")

    effective_weight_unit = weight_unit_override or (weight_unit_col.lower() if weight_unit_col else None)
    if effective_weight_unit not in WEIGHT_UNITS:
        errors.append("Weight Unit must be oz, kg, or lb (or set a batch override)")

    length = _decimal(get(row, "Length"))
    width = _decimal(get(row, "Width"))
    height = _decimal(get(row, "Height"))
    dims_present = [d is not None for d in (length, width, height)]
    effective_dimension_unit = None
    if any(dims_present):
        if not all(dims_present):
            errors.append("Length, Width, and Height must all be present if any is given")
        dimension_unit_col = _cell_str(get(row, "Dimension Unit"))
        effective_dimension_unit = dimension_unit_override or (
            dimension_unit_col.lower() if dimension_unit_col else None
        )
        if effective_dimension_unit not in DIMENSION_UNITS:
            errors.append("Dimension Unit must be inch or cm (or set a batch override)")

    effective_service_code = default_service_code
    if service_code:
        if service_code.upper() not in VALID_SERVICE_CODES:
            errors.append(f"Service Code must be one of {sorted(VALID_SERVICE_CODES)}")
        else:
            effective_service_code = service_code.upper()

    if errors:
        return ParsedRow(row_number=row_number, raw=raw, valid=False, error="; ".join(errors))

    return ParsedRow(
        row_number=row_number,
        raw=raw,
        valid=True,
        label_fields={
            "recipient_name": name,
            "recipient_company": _cell_str(get(row, "Recipient Company")),
            "recipient_address1": address1,
            "recipient_address2": _cell_str(get(row, "Address Line 2")),
            "recipient_city": city,
            "recipient_state": state.upper(),
            "recipient_postal_code": postal_code,
            "recipient_phone": _cell_str(get(row, "Phone")),
            "recipient_email": _cell_str(get(row, "Email")),
            "weight_value": weight_raw,
            "weight_unit": effective_weight_unit,
            "length_value": length,
            "width_value": width,
            "height_value": height,
            "dimension_unit": effective_dimension_unit,
            "declared_value": declared_value,
            "reference1": _cell_str(get(row, "Reference / Order #")),
            "service_code": effective_service_code,
        },
    )
