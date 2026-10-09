"""Bulk run report (.xlsx) and label zip — generated on demand (D18), never stored."""

import io
import zipfile

from openpyxl import Workbook

from app.bulk.template import TEMPLATE_HEADERS
from app.storage import resolve_under_root


def build_report_xlsx(rows: list, labels_by_id: dict) -> bytes:
    wb = Workbook()

    all_sheet = wb.active
    all_sheet.title = "All Rows"
    all_sheet.append(["Row", "Status", "Error", *TEMPLATE_HEADERS])
    for row in rows:
        label = labels_by_id.get(row.label_id)
        error = row.validation_error or row.epg_error_message or ""
        all_sheet.append(
            [row.row_number, row.status, error, *[row.raw.get(h) for h in TEMPLATE_HEADERS]]
        )

    # Only rows where NOTHING was bought: safe to fix and re-upload.
    failed_sheet = wb.create_sheet("Failed Rows")
    failed_sheet.append(TEMPLATE_HEADERS)
    for row in rows:
        if row.status in ("failed", "invalid"):
            failed_sheet.append([row.raw.get(h) for h in TEMPLATE_HEADERS])

    # SHIP-3: rows that may have been bought (needs checking) or were refused as a duplicate. Never re-upload these;
    # check the label in the portal first.
    check_rows = [r for r in rows if r.status in ("needs_checking", "duplicate")]
    if check_rows:
        check_sheet = wb.create_sheet("Needs checking - do NOT re-upload")
        check_sheet.append(["Row", "Status", "Label #", "What happened", *TEMPLATE_HEADERS])
        for row in check_rows:
            check_sheet.append([row.row_number, row.status, row.label_id, row.epg_error_message or "",
                                *[row.raw.get(h) for h in TEMPLATE_HEADERS]])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_labels_zip(labels: list) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for label in labels:
            if not label.pdf_path:
                continue
            path = resolve_under_root(label.pdf_path)
            if path.exists():
                zf.write(path, arcname=path.name)
    return out.getvalue()
